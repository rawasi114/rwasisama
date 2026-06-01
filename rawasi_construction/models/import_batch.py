# -*- coding: utf-8 -*-
"""استيراد جدول الكميات من Excel — رفع، مطابقة آلية (4 مستويات)، مراجعة بشرية، اعتماد.

كل دفعة **يجب** ربطها بمنافسة. عند الاعتماد تتحوّل السطور المراجَعة إلى بنود
rawasi.boq.item على تلك المنافسة، وتعزّز المطابقاتُ المقبولةُ «الصياغات البديلة»،
ويُسجَّل كل قرار في rawasi.import.audit. المطابقة لا تعتمد تلقائياً أبداً — يراجع البشر دائماً.
"""
import base64
import io
from difflib import SequenceMatcher

from odoo import api, fields, models, _
from odoo.exceptions import UserError

from .arabic_utils import normalize_light, normalize_match

# كلمات مفتاحية للعناوين (مطبَّعة) ← العمود المنطقي. أول تطابق يفوز.
_HEADER_MAP = {
    "serial": ["الرقم التسلسلي", "م", "تسلسل", "رقم"],
    "category": ["الفئة", "التصنيف"],
    "description": ["الوصف", "البند", "بيان"],
    "specification": ["المواصفات", "مواصفات"],
    "qty": ["الكمية", "كمية"],
    "unit": ["وحدة القياس", "الوحدة", "وحده"],
    "unit_cost": ["التكلفة الافرادي", "تكلفة الوحدة", "التكلفه الافرادي"],
    "total_cost": ["اجمالي التكلفة للبند", "اجمالي التكلفه"],
    "unit_price": ["السعر الافرادي", "سعر الوحدة"],
    "subtotal": ["اجمالي البند", "الاجمالي", "اجمالي"],
}

FUZZY_THRESHOLD = 0.60


class RawasiImportBatch(models.Model):
    _name = "rawasi.import.batch"
    _description = "دفعة استيراد جدول كميات"
    _inherit = ["mail.thread", "mail.activity.mixin"]
    _order = "upload_date desc, id desc"
    _rec_name = "filename"

    filename = fields.Char(string="اسم الملف", required=True, tracking=True)
    file_data = fields.Binary(string="ملف Excel", attachment=True, required=True)
    upload_date = fields.Datetime(
        string="تاريخ الرفع", default=fields.Datetime.now, readonly=True, index=True
    )
    uploaded_by_user_id = fields.Many2one(
        "res.users", string="رفعه", default=lambda self: self.env.user, readonly=True
    )
    competition_id = fields.Many2one(
        "rawasi.competition", string="المنافسة المرتبطة",
        required=True, ondelete="restrict", index=True, tracking=True,
        help="المنافسة التي ستنتقل إليها بنود جدول الكميات بعد الاعتماد.",
    )
    source_partner_id = fields.Many2one("res.partner", string="الجهة المصدر")
    notes = fields.Text(string="ملاحظات")
    state = fields.Selection(
        [
            ("draft", "مسودة"),
            ("reviewing", "قيد المراجعة"),
            ("approved", "معتمَدة"),
            ("cancelled", "ملغاة"),
        ],
        string="الحالة", default="draft", required=True, tracking=True, index=True,
    )
    line_ids = fields.One2many(
        "rawasi.import.batch.line", "import_batch_id", string="سطور الجدول"
    )
    total_lines = fields.Integer(compute="_compute_counts", store=True)
    matched_lines = fields.Integer(compute="_compute_counts", store=True, string="مطابَقة")
    pending_lines = fields.Integer(compute="_compute_counts", store=True, string="بانتظار المراجعة")

    @api.depends("line_ids", "line_ids.review_status", "line_ids.suggested_match_id")
    def _compute_counts(self):
        for batch in self:
            lines = batch.line_ids
            batch.total_lines = len(lines)
            batch.matched_lines = len(lines.filtered(lambda l: l.suggested_match_id))
            batch.pending_lines = len(lines.filtered(lambda l: l.review_status == "pending"))

    # ------------------------------------------------------------------
    # تحليل ملف Excel
    # ------------------------------------------------------------------
    def action_parse(self):
        self.ensure_one()
        if self.state not in ("draft",):
            raise UserError(_("لا يمكن تحليل الملف إلا في حالة المسودة."))
        self.line_ids.unlink()
        rows = self._read_xlsx()
        if not rows:
            raise UserError(_("الملف فارغ أو غير مقروء."))
        header_index = self._detect_columns(rows)
        if "description" not in header_index:
            raise UserError(_("تعذّر العثور على عمود «الوصف» في الملف."))
        header_row = self._header_row_number(rows, header_index)
        Line = self.env["rawasi.import.batch.line"]
        line_no = 0
        for r_idx, row in enumerate(rows):
            if r_idx <= header_row:
                continue
            desc = self._cell(row, header_index.get("description"))
            if not desc or not str(desc).strip():
                continue
            line_no += 1
            Line.create(self._row_to_line_vals(row, header_index, line_no))
        self.state = "reviewing"
        # تشغيل المطابقة على كل السطور
        self.line_ids._run_matching()
        return self.action_open_review()

    def _read_xlsx(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise UserError(_("مكتبة openpyxl غير متاحة على الخادم."))
        data = base64.b64decode(self.file_data)
        wb = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows

    def _detect_columns(self, rows):
        """يفحص أول 10 صفوف بحثاً عن صف عناوين يطابق الأعمدة المعروفة."""
        best = {}
        for row in rows[:10]:
            mapping = {}
            for col_idx, cell in enumerate(row):
                norm = normalize_match(cell) if cell is not None else ""
                if not norm:
                    continue
                for logical, keywords in _HEADER_MAP.items():
                    if logical in mapping:
                        continue
                    if any(normalize_match(k) == norm for k in keywords):
                        mapping[logical] = col_idx
            if len(mapping) > len(best):
                best = mapping
        return best

    def _header_row_number(self, rows, header_index):
        target_cols = set(header_index.values())
        for r_idx, row in enumerate(rows[:10]):
            row_cols = set()
            for col_idx, cell in enumerate(row):
                if col_idx in target_cols and cell is not None and str(cell).strip():
                    row_cols.add(col_idx)
            if row_cols == target_cols and target_cols:
                return r_idx
        return 0

    @staticmethod
    def _cell(row, idx):
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    def _row_to_line_vals(self, row, header_index, line_no):
        def text(key):
            val = self._cell(row, header_index.get(key))
            return normalize_light(str(val)) if val is not None else False

        def number(key):
            val = self._cell(row, header_index.get(key))
            if val is None:
                return 0.0
            try:
                return float(str(val).replace(",", "").strip())
            except (ValueError, AttributeError):
                return 0.0

        return {
            "import_batch_id": self.id,
            "line_number": line_no,
            "original_serial": text("serial"),
            "original_category": text("category"),
            "original_text": text("description") or "—",
            "original_specification": text("specification"),
            "original_unit": text("unit"),
            "original_qty": number("qty"),
            "original_unit_cost": number("unit_cost"),
            "original_unit_price": number("unit_price"),
        }

    # ------------------------------------------------------------------
    # أزرار سير الحالة
    # ------------------------------------------------------------------
    def action_open_review(self):
        self.ensure_one()
        return {
            "type": "ir.actions.act_window",
            "name": _("مراجعة دفعة: %s") % self.filename,
            "res_model": "rawasi.import.batch.line",
            "view_mode": "list,form",
            "domain": [("import_batch_id", "=", self.id)],
            "context": {"default_import_batch_id": self.id, "search_default_pending": 1},
        }

    def action_cancel(self):
        self.write({"state": "cancelled"})

    def action_rerun_matching(self):
        self.ensure_one()
        self.line_ids._run_matching()

    def action_approve(self):
        """يرحّل السطور المراجَعة إلى rawasi.boq.item على المنافسة المرتبطة."""
        self.ensure_one()
        if self.state in ("approved", "cancelled"):
            raise UserError(_("الدفعة في حالة «%s» — لا يمكن اعتمادها.") % self.state)
        Audit = self.env["rawasi.import.audit"]
        BoqItem = self.env["rawasi.boq.item"]
        Variant = self.env["rawasi.item.variant"]
        created = 0
        for line in self.line_ids:
            status = line.review_status
            # القرارات المعلَّقة: تُحفظ بلا ترشيح (لا اعتماد تلقائي للترشيح)
            if status == "pending":
                status = "saved" if not line.final_match_id else "accepted"
            ref = line.final_match_id or (
                line.suggested_match_id if status == "accepted" else False
            )
            if status == "rejected":
                Audit.create(line._audit_vals("rejected", False))
                continue
            if ref:
                Variant.learn(ref, line.original_text, self.source_partner_id)
                Audit.create(line._audit_vals(
                    "modified" if line.final_match_id and
                    line.final_match_id != line.suggested_match_id else "accepted",
                    ref,
                ))
            else:
                Audit.create(line._audit_vals("saved_no_match", False))
            vals = line._build_boq_values(self.competition_id)
            if ref:
                vals["reference_item_id"] = ref.id
            BoqItem.create(vals)
            created += 1
        self.state = "approved"
        self.message_post(
            body=_("تم اعتماد الاستيراد — أُضيف %d بنداً للمنافسة «%s».")
            % (created, self.competition_id.name)
        )
        return {
            "type": "ir.actions.act_window",
            "res_model": "rawasi.competition",
            "res_id": self.competition_id.id,
            "view_mode": "form",
            "target": "current",
        }


class RawasiImportBatchLine(models.Model):
    _name = "rawasi.import.batch.line"
    _description = "سطر استيراد جدول كميات"
    _order = "import_batch_id, line_number"

    import_batch_id = fields.Many2one(
        "rawasi.import.batch", string="الدفعة", required=True, ondelete="cascade", index=True
    )
    line_number = fields.Integer(string="رقم السطر", required=True)
    # الأعمدة كما وردت
    original_serial = fields.Char(string="الرقم التسلسلي")
    original_category = fields.Char(string="الفئة")
    original_text = fields.Text(string="الوصف", required=True)
    original_specification = fields.Text(string="المواصفات")
    original_unit = fields.Char(string="وحدة القياس")
    original_qty = fields.Float(string="الكمية")
    original_unit_cost = fields.Float(string="التكلفة الإفرادي")
    original_unit_price = fields.Float(string="السعر الإفرادي")
    # ناتج المطابقة
    suggested_match_id = fields.Many2one("rawasi.reference.item", string="الترشيح المقترح")
    suggestion_confidence = fields.Float(string="درجة الثقة %")
    suggestion_source = fields.Selection(
        [
            ("exact_text", "تطابق نصي كامل"),
            ("lcgpa_code", "تطابق برمز LCGPA"),
            ("fuzzy", "تطابق تقريبي"),
            ("none", "بلا ترشيح"),
        ],
        string="مصدر الترشيح", default="none",
    )
    detected_lcgpa_id = fields.Many2one("rawasi.lcgpa.code", string="رمز LCGPA المكتشف")
    # المراجعة
    review_status = fields.Selection(
        [
            ("pending", "بانتظار المراجعة"),
            ("accepted", "قُبل الترشيح"),
            ("modified", "عُدِّل الترشيح"),
            ("saved", "محفوظ بلا ترشيح"),
            ("rejected", "مرفوض"),
        ],
        string="حالة المراجعة", default="pending", required=True, index=True,
    )
    final_match_id = fields.Many2one("rawasi.reference.item", string="البند المرجعي النهائي")
    reviewer_notes = fields.Text(string="ملاحظات المراجع")

    # ------------------------------------------------------------------
    # محرك المطابقة — 4 مستويات، لا اعتماد تلقائي
    # ------------------------------------------------------------------
    def _run_matching(self):
        Variant = self.env["rawasi.item.variant"]
        RefItem = self.env["rawasi.reference.item"]
        Lcgpa = self.env["rawasi.lcgpa.code"]
        ref_cache = RefItem.search([])
        ref_norms = {ref: normalize_match(ref.name) for ref in ref_cache}
        for line in self:
            norm = normalize_match(line.original_text)
            vals = {
                "suggested_match_id": False,
                "suggestion_confidence": 0.0,
                "suggestion_source": "none",
                "detected_lcgpa_id": False,
            }
            if not norm:
                line.write(vals)
                continue
            # المستوى 1: تطابق نصي كامل عبر الصياغات البديلة ثم أسماء البنود
            variant = Variant.search([("normalized_text", "=", norm)], limit=1)
            match, source, conf = False, "none", 0.0
            if variant:
                match, source, conf = variant.reference_item_id, "exact_text", 100.0
            elif norm in ref_norms.values():
                match = next(r for r, n in ref_norms.items() if n == norm)
                source, conf = "exact_text", 95.0
            # المستوى 2: تطابق برمز LCGPA
            if not match and line.original_category:
                lcgpa = Lcgpa.match_code(line.original_category)
                if lcgpa and lcgpa.reference_item_ids:
                    vals["detected_lcgpa_id"] = lcgpa.id
                    match, source, conf = lcgpa.reference_item_ids[0], "lcgpa_code", 80.0
            # المستوى 3: تطابق تقريبي
            if not match:
                best_ref, best_score = False, 0.0
                matcher = SequenceMatcher()
                matcher.set_seq2(norm)
                for ref, rnorm in ref_norms.items():
                    matcher.set_seq1(rnorm)
                    score = matcher.ratio()
                    if score > best_score:
                        best_ref, best_score = ref, score
                if best_ref and best_score >= FUZZY_THRESHOLD:
                    match, source, conf = best_ref, "fuzzy", round(best_score * 100, 1)
            vals.update(
                suggested_match_id=match.id if match else False,
                suggestion_source=source,
                suggestion_confidence=conf,
            )
            line.write(vals)

    def action_accept(self):
        for line in self:
            if not line.suggested_match_id:
                raise UserError(_("لا يوجد ترشيح لقبوله. عدِّل البند أو ارفضه."))
            line.write(
                {"review_status": "accepted", "final_match_id": line.suggested_match_id.id}
            )

    def action_reject(self):
        self.write({"review_status": "rejected", "final_match_id": False})

    @api.onchange("final_match_id")
    def _onchange_final_match(self):
        if self.final_match_id and self.final_match_id != self.suggested_match_id:
            self.review_status = "modified"

    def _resolve_uom(self):
        """يحاول مطابقة نص الوحدة الخام بـ uom.uom؛ وإلا يرجع لوحدة «عدد»."""
        self.ensure_one()
        Uom = self.env["uom.uom"]
        if self.original_unit:
            norm = normalize_match(self.original_unit)
            for uom in Uom.search([]):
                if normalize_match(uom.name) == norm:
                    return uom
        return self.env.ref("uom.product_uom_unit", raise_if_not_found=False) or Uom.browse()

    def _build_boq_values(self, competition):
        self.ensure_one()
        uom = self._resolve_uom()
        return {
            "competition_id": competition.id,
            "sequence": self.line_number * 10,
            "serial": self.original_serial or str(self.line_number),
            "category": self.original_category or False,
            "description": self.original_text or "—",
            "specification": self.original_specification or False,
            "uom_id": uom.id if uom else False,
            "qty": self.original_qty or 0.0,
            "unit_cost": self.original_unit_cost or 0.0,
            "unit_price": self.original_unit_price or 0.0,
        }

    def _audit_vals(self, user_action, ref):
        self.ensure_one()
        return {
            "original_text": self.original_text,
            "reference_item_id": ref.id if ref else False,
            "suggested_item_id": self.suggested_match_id.id or False,
            "match_source": self.suggestion_source,
            "confidence_score": self.suggestion_confidence,
            "user_action": user_action,
            "import_batch_id": self.import_batch_id.id,
            "competition_id": self.import_batch_id.competition_id.id,
        }
