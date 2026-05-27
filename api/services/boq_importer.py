"""BoQ file import service.

Orchestrates the full import pipeline for an uploaded BoQ file:

1. Persist the original file under tender_files (audit trail).
2. Parse the file:
   - Excel/CSV → local parser (fast, free).
   - PDF → Claude API extractor (slower, costs tokens).
3. Normalize each extracted item against the master items dictionary.
4. Stage the items for human review; the caller confirms before they
   are committed to ``tender_boq_items`` or ``rawasi_bid_items``.

The service is split from the FastAPI router so it can be unit-tested
without HTTP plumbing.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Literal
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from api.core.errors import NotFoundError, ValidationError
from api.core.logging import get_logger
from api.models import (
    Competitor,
    RawasiBid,
    RawasiBidItem,
    Tender,
    TenderBoqItem,
    TenderFile,
)
from api.services.boq_normalizer import BoqNormalizer

LOG = get_logger(__name__)

BoqType = Literal["original_tender", "rawasi", "competitor"]
VALID_BOQ_TYPES: tuple[BoqType, ...] = ("original_tender", "rawasi", "competitor")

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024
SUPPORTED_EXTENSIONS: dict[str, str] = {
    ".xlsx": "excel",
    ".xls": "excel",
    ".csv": "csv",
    ".pdf": "pdf",
}
SUPPORTED_MIME_TYPES: dict[str, str] = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "excel",
    "application/vnd.ms-excel": "excel",
    "text/csv": "csv",
    "application/pdf": "pdf",
}


@dataclass
class ParsedBoqItem:
    """One row parsed from a BoQ file, before normalization."""

    sequence_number: str | None
    description: str
    unit: str | None
    quantity: float
    unit_price: float | None = None
    total_price: float | None = None


@dataclass
class StagedBoqItem:
    """A parsed item enriched with normalization metadata, ready for review."""

    sequence_number: str | None
    description: str
    unit: str | None
    quantity: float
    unit_price: float | None
    total_price: float | None
    master_item_id: UUID | None
    master_item_code: str | None
    matching_confidence: float
    matching_method: str
    requires_manual_review: bool


@dataclass
class ImportPreview:
    """The full result returned to the caller for review."""

    tender_id: UUID
    boq_type: BoqType
    source_file_id: UUID
    source_competitor_id: UUID | None
    items: list[StagedBoqItem] = field(default_factory=list)
    parser: str = ""
    warnings: list[str] = field(default_factory=list)
    total_items: int = 0
    items_needing_review: int = 0


class BoqImporter:
    """Parses BoQ files and stages items for confirmation.

    Storage strategy:
    - Original files saved under ``settings.uploads_root / <tender_id>/boq/``.
    - Parsed items kept in memory inside ``ImportPreview`` until ``confirm``
      is called — they are NOT written to ``tender_boq_items`` until then.
    """

    def __init__(
        self,
        db: Session,
        normalizer: BoqNormalizer | None = None,
        claude_extractor=None,
        uploads_root: Path | None = None,
    ) -> None:
        self.db = db
        self.normalizer = normalizer or BoqNormalizer(db)
        self.claude_extractor = claude_extractor
        self.uploads_root = uploads_root or Path("uploads")

    def import_file(
        self,
        tender_id: UUID,
        *,
        file_bytes: bytes,
        filename: str,
        mime_type: str | None,
        boq_type: BoqType,
        source_competitor_id: UUID | None = None,
    ) -> ImportPreview:
        """Parse + normalize + stage. Does NOT commit BoQ items."""
        self._validate_boq_type(boq_type)
        self._validate_file_size(file_bytes)
        parser = self._detect_parser(filename, mime_type)
        tender = self._require_tender(tender_id)
        self._validate_competitor_for_type(boq_type, source_competitor_id)

        stored_file = self._persist_file(
            tender, filename, file_bytes, mime_type, boq_type
        )

        if parser == "excel":
            parsed = self._parse_excel(file_bytes)
        elif parser == "csv":
            parsed = self._parse_csv(file_bytes)
        elif parser == "pdf":
            parsed = self._parse_pdf(file_bytes, mime_type or "application/pdf")
        else:  # pragma: no cover — guarded by _detect_parser
            raise ValidationError(f"unsupported parser: {parser}")

        if not parsed:
            raise ValidationError(
                "لم يتم استخراج أي بنود من الملف. تحقق من بنية الملف."
            )

        staged: list[StagedBoqItem] = []
        warnings: list[str] = []
        for item in parsed:
            try:
                outcome = self.normalizer.normalize(
                    item.description, item.unit, use_claude=False
                )
                staged.append(
                    StagedBoqItem(
                        sequence_number=item.sequence_number,
                        description=item.description,
                        unit=item.unit,
                        quantity=item.quantity,
                        unit_price=item.unit_price,
                        total_price=item.total_price,
                        master_item_id=outcome.master_item_id,
                        master_item_code=outcome.master_item_code,
                        matching_confidence=float(outcome.confidence or 0.0),
                        matching_method=outcome.method,
                        requires_manual_review=outcome.requires_manual_review,
                    )
                )
            except Exception as exc:  # pragma: no cover — defensive
                warnings.append(f"row '{item.description[:50]}': {exc}")

        items_needing_review = sum(1 for s in staged if s.requires_manual_review)
        self.db.commit()

        return ImportPreview(
            tender_id=tender_id,
            boq_type=boq_type,
            source_file_id=stored_file.id,
            source_competitor_id=source_competitor_id,
            items=staged,
            parser=parser,
            warnings=warnings,
            total_items=len(staged),
            items_needing_review=items_needing_review,
        )

    def confirm(
        self,
        tender_id: UUID,
        boq_type: BoqType,
        items: list[StagedBoqItem],
        source_file_id: UUID | None = None,
        source_competitor_id: UUID | None = None,
    ) -> int:
        """Commit the staged items to the correct table for ``boq_type``."""
        self._validate_boq_type(boq_type)
        tender = self._require_tender(tender_id)
        self._validate_competitor_for_type(boq_type, source_competitor_id)

        if boq_type == "rawasi":
            count = self._commit_rawasi_items(tender, items)
        else:
            count = self._commit_tender_boq_items(
                tender, boq_type, items, source_file_id, source_competitor_id
            )

        tender.has_full_boq = True
        self.db.commit()
        return count

    # ------------------------------------------------------------------ #
    # Validation                                                         #
    # ------------------------------------------------------------------ #

    @staticmethod
    def _validate_boq_type(boq_type: str) -> None:
        if boq_type not in VALID_BOQ_TYPES:
            raise ValidationError(
                f"boq_type must be one of {VALID_BOQ_TYPES}, got {boq_type!r}"
            )

    @staticmethod
    def _validate_file_size(file_bytes: bytes) -> None:
        if len(file_bytes) > MAX_FILE_SIZE_BYTES:
            raise ValidationError(
                f"حجم الملف {len(file_bytes) / 1024 / 1024:.1f}MB "
                f"يتجاوز الحد الأقصى ({MAX_FILE_SIZE_BYTES // 1024 // 1024}MB)"
            )
        if not file_bytes:
            raise ValidationError("الملف فارغ")

    @staticmethod
    def _detect_parser(filename: str, mime_type: str | None) -> str:
        if mime_type and mime_type in SUPPORTED_MIME_TYPES:
            return SUPPORTED_MIME_TYPES[mime_type]
        ext = Path(filename).suffix.lower()
        if ext in SUPPORTED_EXTENSIONS:
            return SUPPORTED_EXTENSIONS[ext]
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS.keys()))
        raise ValidationError(
            f"صيغة الملف غير مدعومة. الصيغ المدعومة: {supported}"
        )

    def _require_tender(self, tender_id: UUID) -> Tender:
        tender = self.db.get(Tender, tender_id)
        if tender is None:
            raise NotFoundError(f"tender {tender_id} not found")
        return tender

    def _validate_competitor_for_type(
        self, boq_type: BoqType, competitor_id: UUID | None
    ) -> None:
        if boq_type == "competitor":
            if competitor_id is None:
                raise ValidationError(
                    "boq_type=competitor requires source_competitor_id"
                )
            if self.db.get(Competitor, competitor_id) is None:
                raise ValidationError(
                    f"competitor {competitor_id} not found"
                )
        elif competitor_id is not None:
            raise ValidationError(
                "source_competitor_id is only valid when boq_type=competitor"
            )

    # ------------------------------------------------------------------ #
    # Storage                                                            #
    # ------------------------------------------------------------------ #

    def _persist_file(
        self,
        tender: Tender,
        filename: str,
        file_bytes: bytes,
        mime_type: str | None,
        boq_type: BoqType,
    ) -> TenderFile:
        digest = hashlib.sha256(file_bytes).hexdigest()[:16]
        safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", filename)
        timestamp = datetime.utcnow().strftime("%Y%m%dT%H%M%S")
        rel_path = (
            Path(str(tender.id)) / "boq" / f"{timestamp}_{digest}_{safe_name}"
        )
        abs_path = self.uploads_root / rel_path
        abs_path.parent.mkdir(parents=True, exist_ok=True)
        abs_path.write_bytes(file_bytes)

        record = TenderFile(
            tender_id=tender.id,
            file_type=f"boq_{boq_type}",
            filename=filename,
            storage_path=str(rel_path),
            mime_type=mime_type,
            size_bytes=len(file_bytes),
        )
        self.db.add(record)
        self.db.flush()
        return record

    # ------------------------------------------------------------------ #
    # Parsers                                                            #
    # ------------------------------------------------------------------ #

    def _parse_excel(self, file_bytes: bytes) -> list[ParsedBoqItem]:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        items: list[ParsedBoqItem] = []
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            header_row, header_idx = self._detect_header(rows)
            if header_row is None or header_idx is None:
                continue
            mapping = self._map_columns(header_row)
            if "description" not in mapping or "quantity" not in mapping:
                continue
            for row in rows[header_idx + 1 :]:
                item = self._row_to_item(row, mapping)
                if item is not None:
                    items.append(item)
        return items

    def _parse_csv(self, file_bytes: bytes) -> list[ParsedBoqItem]:
        text = self._decode_text(file_bytes)
        reader = csv.reader(io.StringIO(text))
        rows = [tuple(r) for r in reader]
        if not rows:
            return []
        header_row, header_idx = self._detect_header(rows)
        if header_row is None or header_idx is None:
            return []
        mapping = self._map_columns(header_row)
        if "description" not in mapping or "quantity" not in mapping:
            return []
        items: list[ParsedBoqItem] = []
        for row in rows[header_idx + 1 :]:
            item = self._row_to_item(tuple(row), mapping)
            if item is not None:
                items.append(item)
        return items

    def _parse_pdf(self, file_bytes: bytes, mime_type: str) -> list[ParsedBoqItem]:
        if self.claude_extractor is None:
            raise ValidationError(
                "PDF parsing requires Claude API but no extractor is configured"
            )
        response = self.claude_extractor.extract_boq(file_bytes, mime_type)
        raw_items = response.get("boq_items", [])
        items: list[ParsedBoqItem] = []
        for entry in raw_items:
            qty = self._coerce_float(entry.get("quantity"))
            if qty is None or qty <= 0:
                continue
            items.append(
                ParsedBoqItem(
                    sequence_number=str(entry.get("sequence_number") or "") or None,
                    description=str(entry.get("description_ar") or "").strip(),
                    unit=str(entry.get("unit") or "") or None,
                    quantity=qty,
                    unit_price=self._coerce_float(entry.get("estimated_unit_price")),
                )
            )
        return items

    # ------------------------------------------------------------------ #
    # Excel / CSV helpers                                                #
    # ------------------------------------------------------------------ #

    HEADER_KEYWORDS: dict[str, tuple[str, ...]] = {
        "sequence_number": ("seq", "no", "رقم", "م", "م.", "تسلسل", "بند"),
        "description": ("description", "desc", "البيان", "الوصف", "البند", "اسم"),
        "unit": ("unit", "uom", "الوحدة", "وحدة"),
        "quantity": ("qty", "quantity", "الكمية", "كمية"),
        "unit_price": (
            "unit price",
            "price",
            "rate",
            "سعر الوحدة",
            "السعر",
            "سعر",
        ),
        "total_price": (
            "total",
            "amount",
            "الإجمالي",
            "اجمالي",
            "المجموع",
            "القيمة",
        ),
    }

    def _detect_header(
        self, rows: list[tuple]
    ) -> tuple[tuple | None, int | None]:
        for idx, row in enumerate(rows[:20]):
            normalized = tuple(self._cell_text(c).lower() for c in row)
            if any(
                any(kw in cell for kw in self.HEADER_KEYWORDS["description"])
                for cell in normalized
                if cell
            ) and any(
                any(kw in cell for kw in self.HEADER_KEYWORDS["quantity"])
                for cell in normalized
                if cell
            ):
                return row, idx
        return None, None

    def _map_columns(self, header_row: tuple) -> dict[str, int]:
        mapping: dict[str, int] = {}
        for idx, cell in enumerate(header_row):
            text = self._cell_text(cell).lower()
            if not text:
                continue
            for field_name, keywords in self.HEADER_KEYWORDS.items():
                if field_name in mapping:
                    continue
                if any(kw in text for kw in keywords):
                    mapping[field_name] = idx
                    break
        return mapping

    def _row_to_item(
        self, row: tuple, mapping: dict[str, int]
    ) -> ParsedBoqItem | None:
        if not row:
            return None
        description = self._cell_text(row[mapping["description"]]).strip()
        if not description:
            return None
        quantity = self._coerce_float(row[mapping["quantity"]])
        if quantity is None or quantity <= 0:
            return None
        return ParsedBoqItem(
            sequence_number=(
                self._cell_text(row[mapping["sequence_number"]]).strip() or None
                if "sequence_number" in mapping
                else None
            ),
            description=description,
            unit=(
                self._cell_text(row[mapping["unit"]]).strip() or None
                if "unit" in mapping
                else None
            ),
            quantity=quantity,
            unit_price=(
                self._coerce_float(row[mapping["unit_price"]])
                if "unit_price" in mapping
                else None
            ),
            total_price=(
                self._coerce_float(row[mapping["total_price"]])
                if "total_price" in mapping
                else None
            ),
        )

    @staticmethod
    def _cell_text(cell) -> str:
        if cell is None:
            return ""
        return str(cell)

    @staticmethod
    def _coerce_float(value) -> float | None:
        if value is None or value == "":
            return None
        try:
            return float(value)
        except (TypeError, ValueError):
            try:
                return float(Decimal(str(value).replace(",", "")))
            except (InvalidOperation, ValueError):
                return None

    @staticmethod
    def _decode_text(file_bytes: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "cp1256", "windows-1256"):
            try:
                return file_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return file_bytes.decode("utf-8", errors="replace")

    # ------------------------------------------------------------------ #
    # Commit                                                             #
    # ------------------------------------------------------------------ #

    def _commit_tender_boq_items(
        self,
        tender: Tender,
        boq_type: BoqType,
        items: list[StagedBoqItem],
        source_file_id: UUID | None,
        source_competitor_id: UUID | None,
    ) -> int:
        for item in items:
            row = TenderBoqItem(
                tender_id=tender.id,
                sequence_number=item.sequence_number,
                original_description=item.description,
                original_unit=item.unit,
                quantity=item.quantity,
                unit_estimated_price=item.unit_price,
                master_item_id=item.master_item_id,
                matching_confidence=item.matching_confidence,
                matching_method=item.matching_method,
                matched_at=datetime.utcnow() if item.master_item_id else None,
                requires_manual_review=item.requires_manual_review,
                boq_type=boq_type,
                source_file_id=source_file_id,
                source_competitor_id=source_competitor_id,
            )
            self.db.add(row)
        self.db.flush()
        return len(items)

    def _commit_rawasi_items(
        self, tender: Tender, items: list[StagedBoqItem]
    ) -> int:
        bid = self._get_or_create_rawasi_bid(tender, items)
        for item in items:
            unit_price = item.unit_price or 0.0
            total_price = item.total_price or unit_price * item.quantity
            self.db.add(
                RawasiBidItem(
                    rawasi_bid_id=bid.id,
                    master_item_id=item.master_item_id,
                    quantity=item.quantity,
                    unit_price=unit_price,
                    total_price=total_price,
                )
            )
        self.db.flush()
        return len(items)

    def _get_or_create_rawasi_bid(
        self, tender: Tender, items: list[StagedBoqItem]
    ) -> RawasiBid:
        bid = self.db.query(RawasiBid).filter_by(tender_id=tender.id).one_or_none()
        if bid is not None:
            return bid
        total = sum(
            (item.total_price or (item.unit_price or 0) * item.quantity)
            for item in items
        )
        bid = RawasiBid(
            tender_id=tender.id,
            bid_total_amount=total,
            won=tender.awarded_to_competitor_id is None,
        )
        self.db.add(bid)
        self.db.flush()
        tender.rawasi_participated = True
        tender.rawasi_bid_id = str(bid.id)
        return bid


__all__ = [
    "BoqImporter",
    "BoqType",
    "ImportPreview",
    "ParsedBoqItem",
    "StagedBoqItem",
    "VALID_BOQ_TYPES",
]
