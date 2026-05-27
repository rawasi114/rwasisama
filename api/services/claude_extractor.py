"""Claude API extraction service.

Wraps Anthropic Claude calls to extract structured data from PDFs and Excel:
- Award announcements
- BoQ tables
- Item normalization assists

The service is kept dependency-light at construction time so it can be
instantiated even when the API key is missing (for unit testing). Methods
will raise :class:`ClaudeExtractionError` if called without a configured
client.
"""

from __future__ import annotations

import base64
import json
from pathlib import Path
from typing import Any

from api.core.config import Settings, get_settings
from api.core.errors import ClaudeExtractionError
from api.core.logging import get_logger

LOG = get_logger(__name__)
PROMPTS_DIR = Path(__file__).resolve().parent.parent.parent / "prompts"


def _load_prompt(name: str) -> str:
    path = PROMPTS_DIR / name
    if not path.exists():
        raise FileNotFoundError(f"Prompt file not found: {path}")
    return path.read_text(encoding="utf-8")


class ClaudeExtractor:
    """Wraps Claude API calls for the system's three extraction tasks."""

    def __init__(self, settings: Settings | None = None, client: Any = None) -> None:
        self.settings = settings or get_settings()
        self._client = client

    @property
    def client(self):
        if self._client is None:
            try:
                from anthropic import Anthropic
            except ImportError as exc:  # pragma: no cover
                raise ClaudeExtractionError(
                    "anthropic SDK not installed"
                ) from exc
            if not self.settings.anthropic_api_key:
                raise ClaudeExtractionError(
                    "ANTHROPIC_API_KEY is not configured"
                )
            self._client = Anthropic(api_key=self.settings.anthropic_api_key)
        return self._client

    def extract_award_announcement(self, pdf_bytes: bytes) -> dict:
        """Parse an award-announcement PDF into a structured dict."""
        prompt = _load_prompt("extract_award_announcement.md")
        return self._call_with_document(pdf_bytes, "application/pdf", prompt)

    def extract_boq(self, file_bytes: bytes, mime_type: str) -> dict:
        """Parse a BoQ document into structured items."""
        prompt = _load_prompt("extract_boq.md")
        if mime_type == "application/pdf":
            return self._call_with_document(file_bytes, mime_type, prompt)
        text_content = self._excel_to_text(file_bytes)
        combined = f"{prompt}\n\n# جدول الكميات:\n{text_content}"
        return self._call_text_only(combined)

    def normalize_item(
        self,
        original_description: str,
        original_unit: str,
        candidates: list[dict],
    ) -> dict:
        """Ask Claude to pick the best matching master item."""
        template = _load_prompt("normalize_boq_item.md")
        candidates_text = "\n".join(
            f"- {c['code']} | {c['name_ar']} | "
            f"{c.get('description_ar', '') or ''} | {c['default_unit']}"
            for c in candidates
        )
        prompt = (
            template.replace("{original_description}", original_description)
            .replace("{original_unit}", original_unit or "")
            .replace("{candidates}", candidates_text)
        )
        return self._call_text_only(prompt, max_tokens=1024)

    def _call_with_document(self, file_bytes: bytes, mime_type: str, prompt: str) -> dict:
        b64 = base64.b64encode(file_bytes).decode("utf-8")
        response = self.client.messages.create(
            model=self.settings.claude_model,
            max_tokens=self.settings.claude_max_tokens,
            messages=[
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "document",
                            "source": {
                                "type": "base64",
                                "media_type": mime_type,
                                "data": b64,
                            },
                        },
                        {"type": "text", "text": prompt},
                    ],
                }
            ],
        )
        return self._parse_json_response(response.content[0].text)

    def _call_text_only(self, prompt: str, *, max_tokens: int | None = None) -> dict:
        response = self.client.messages.create(
            model=self.settings.claude_model,
            max_tokens=max_tokens or self.settings.claude_max_tokens,
            messages=[{"role": "user", "content": prompt}],
        )
        return self._parse_json_response(response.content[0].text)

    @staticmethod
    def _parse_json_response(text: str) -> dict:
        cleaned = text.strip()
        if cleaned.startswith("```json"):
            cleaned = cleaned[7:]
        elif cleaned.startswith("```"):
            cleaned = cleaned[3:]
        if cleaned.endswith("```"):
            cleaned = cleaned[:-3]
        try:
            return json.loads(cleaned.strip())
        except json.JSONDecodeError as exc:
            LOG.error("claude_invalid_json", error=str(exc), preview=cleaned[:200])
            raise ClaudeExtractionError(f"Invalid JSON from Claude: {exc}") from exc

    @staticmethod
    def _excel_to_text(file_bytes: bytes) -> str:
        """Convert an Excel workbook into a plain-text table the model can read."""
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
        lines: list[str] = []
        for sheet in wb.worksheets:
            lines.append(f"## Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                row_cells = ["" if c is None else str(c) for c in row]
                if any(cell.strip() for cell in row_cells):
                    lines.append(" | ".join(row_cells))
            lines.append("")
        return "\n".join(lines)
