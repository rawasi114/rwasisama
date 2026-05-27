# -*- coding: utf-8 -*-
import base64
import io
from datetime import date

from odoo.tests.common import TransactionCase, tagged

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


def _xlsx(headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue())


@tagged("post_install", "-at_install")
class TestPricedImport(TransactionCase):
    def setUp(self):
        super().setUp()
        self.comp = self.env["rawasi.competition"].create({"name": "منافسة استيراد مُسعّر"})

    def _import(self, headers, rows):
        wiz = self.env["rawasi.boq.import.wizard"].create({
            "competition_id": self.comp.id,
            "file": _xlsx(headers, rows),
            "filename": "boq.xlsx",
        })
        wiz.action_import()

    def test_import_unit_price(self):
        self._import(
            ["وصف البند", "الكمية", "الوحدة", "سعر الوحدة"],
            [["خرسانة مسلحة", 10, "م3", 150], ["حديد تسليح", 5, "طن", 2000]],
        )
        items = self.comp.boq_item_ids
        self.assertEqual(len(items), 2)
        khorsana = items.filtered(lambda i: "خرسانة" in i.name)
        self.assertEqual(khorsana.unit_price, 150.0)
        self.assertEqual(khorsana.total_price, 1500.0)
        self.assertEqual(self.comp.amount_total_price, 11500.0)

    def test_import_derives_unit_price_from_total(self):
        # ملف فيه إجمالي السعر فقط (دون سعر الوحدة) -> يُشتقّ سعر الوحدة
        self._import(
            ["وصف البند", "الكمية", "إجمالي السعر"],
            [["دهان", 20, 1000]],
        )
        item = self.comp.boq_item_ids
        self.assertEqual(item.unit_price, 50.0)  # 1000 / 20
        self.assertEqual(item.total_price, 1000.0)


@tagged("post_install", "-at_install")
class TestProjectsSystem(TransactionCase):
    def test_convert_flags_construction_and_budget(self):
        comp = self.env["rawasi.competition"].create({"name": "منافسة المشروع"})
        self.env["rawasi.boq.item"].create({
            "competition_id": comp.id, "name": "بند", "quantity": 10.0,
            "unit_cost": 100.0, "unit_price": 120.0,
        })
        comp.action_start_pricing()
        comp.action_submit()
        comp.action_won()
        comp.action_convert_to_project()
        proj = comp.project_id
        self.assertTrue(proj)
        self.assertTrue(proj.rawasi_is_construction)
        self.assertEqual(proj.rawasi_competition_id, comp)
        self.assertEqual(proj.rawasi_budget_total, 1000.0)  # 10 × 100
        # يظهر ضمن نطاق قائمة المشاريع
        found = self.env["project.project"].search([
            ("rawasi_is_construction", "=", True), ("id", "=", proj.id),
        ])
        self.assertTrue(found)
        # تم توليد المراحل الخمس
        self.assertTrue(len(proj.wbs_activity_ids) >= 5)


@tagged("post_install", "-at_install")
class TestAwardValue(TransactionCase):
    def setUp(self):
        super().setUp()
        self.comp = self.env["rawasi.competition"].create({"name": "م. ترسية"})
        self.env["rawasi.boq.item"].create({
            "competition_id": self.comp.id, "name": "بند",
            "quantity": 1.0, "unit_cost": 100.0, "unit_price": 120.0,
        })
        self.comp.action_start_pricing()
        self.comp.action_submit()

    def test_close_lost_requires_award_value(self):
        from odoo.exceptions import UserError
        with self.assertRaises(UserError):
            self.comp.action_lost()
        self.comp.award_value = 110.0
        self.comp.awarded_to = "مقاول آخر"
        self.comp.action_lost()
        self.assertEqual(self.comp.state, "lost")
        self.assertEqual(self.comp.award_value, 110.0)

    def test_close_won_autofills_award_value(self):
        # عند الفوز: قيمة الترسية = إجمالي عرضنا شامل الضريبة (إن لم تكن مُحدَّدة)
        self.assertFalse(self.comp.award_value)
        expected = self.comp.amount_total_incl_tax
        self.comp.action_won()
        self.assertEqual(self.comp.state, "won")
        self.assertEqual(self.comp.award_value, expected)


@tagged("post_install", "-at_install")
class TestSerials(TransactionCase):
    def test_competition_serial_and_search(self):
        comp = self.env["rawasi.competition"].create({"name": "منافسة الترميم"})
        self.assertTrue(comp.serial, "يجب توليد رقم تسلسلي")
        self.assertTrue(comp.serial.startswith("RS-TND-"))
        # البحث بالرقم التسلسلي يجد المنافسة
        found = self.env["rawasi.competition"].name_search(name=comp.serial)
        self.assertIn(comp.id, [r[0] for r in found])
        # البحث بالاسم أيضاً
        found_name = self.env["rawasi.competition"].name_search(name="الترميم")
        self.assertIn(comp.id, [r[0] for r in found_name])

    def test_project_serial_on_convert_and_search(self):
        comp = self.env["rawasi.competition"].create({"name": "منافسة المشروع التسلسلي"})
        self.env["rawasi.boq.item"].create({
            "competition_id": comp.id, "name": "بند", "quantity": 1.0,
            "unit_cost": 1.0, "unit_price": 2.0,
        })
        comp.action_start_pricing()
        comp.action_submit()
        comp.action_won()
        comp.action_convert_to_project()
        proj = comp.project_id
        self.assertTrue(proj.rawasi_serial.startswith("RS-PRJ-"))
        found = self.env["project.project"].name_search(name=proj.rawasi_serial)
        self.assertIn(proj.id, [r[0] for r in found])


@tagged("post_install", "-at_install")
class TestScheduleTemplate(TransactionCase):
    def setUp(self):
        super().setUp()
        self.project = self.env["project.project"].create({"name": "مشروع الجدول"})
        self.Wizard = self.env["rawasi.schedule.import.wizard"]

    def test_template_is_valid_xlsx(self):
        data = self.Wizard._generate_schedule_template()
        self.assertTrue(data[:2] == b"PK")  # ملف xlsx صالح
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertIn("الجدول الزمني", wb.sheetnames)
        self.assertIn("المراحل المتاحة", wb.sheetnames)

    def test_download_template_action(self):
        action = self.Wizard.action_download_template()
        self.assertEqual(action["type"], "ir.actions.act_url")
        self.assertIn("/web/content/", action["url"])

    def test_import_all_schedule_fields(self):
        headers = [
            "النشاط", "المرحلة", "التسلسل", "تاريخ البداية", "تاريخ النهاية",
            "بداية خط الأساس", "نهاية خط الأساس", "البداية الفعلية",
            "النهاية الفعلية", "نسبة الإنجاز", "معلم",
        ]
        row = [
            "حفر وردم", "الأعمال الإنشائية", 20,
            "2026-02-01", "2026-02-10",
            "2026-02-01", "2026-02-10",
            "2026-02-02", "2026-02-11",
            "50", "لا",
        ]
        wiz = self.Wizard.create({
            "project_id": self.project.id,
            "file": _xlsx(headers, [row]),
            "filename": "sched.xlsx",
        })
        wiz.action_import()
        act = self.project.wbs_activity_ids
        self.assertEqual(len(act), 1)
        self.assertEqual(act.sequence, 20)
        self.assertEqual(act.date_start, date(2026, 2, 1))
        self.assertEqual(act.date_end, date(2026, 2, 10))
        self.assertEqual(act.baseline_start, date(2026, 2, 1))
        self.assertEqual(act.actual_start, date(2026, 2, 2))
        self.assertEqual(act.actual_end, date(2026, 2, 11))
        self.assertEqual(act.progress, 50.0)
        self.assertFalse(act.is_milestone)
