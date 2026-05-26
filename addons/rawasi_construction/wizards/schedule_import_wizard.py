# -*- coding: utf-8 -*-
"""استيراد الجدول الزمني (WBS) من ملف Excel — مطابقة بالعناوين."""
import base64
import io
from datetime import date, datetime

from odoo import fields, models, _
from odoo.exceptions import UserError

from ..models.arabic_utils import normalize_light

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    openpyxl = None

HEADER_ALIASES = {
    "النشاط": "name", "المهمة": "name", "اسم النشاط": "name", "البند": "name",
    "المرحلة": "phase", "مرحلة": "phase",
    "التسلسل": "sequence", "الترتيب": "sequence", "تسلسل": "sequence", "ترتيب": "sequence",
    "تاريخ البداية": "date_start", "البداية": "date_start", "بداية": "date_start",
    "تاريخ النهاية": "date_end", "النهاية": "date_end", "نهاية": "date_end",
    "بداية خط الأساس": "baseline_start", "بداية الأساس": "baseline_start",
    "نهاية خط الأساس": "baseline_end", "نهاية الأساس": "baseline_end",
    "البداية الفعلية": "actual_start", "بداية فعلية": "actual_start",
    "النهاية الفعلية": "actual_end", "نهاية فعلية": "actual_end",
    "نسبة الإنجاز": "progress", "الإنجاز": "progress", "النسبة": "progress",
    "معلم": "milestone", "معلم رئيسي": "milestone", "معلَم": "milestone",
}
HEADER_SCAN_ROWS = 20

# أعمدة قالب الجدول الزمني (بالترتيب) — تطابق كل حقول الجدول القابلة للاستيراد
TEMPLATE_COLUMNS = [
    "النشاط", "المرحلة", "التسلسل",
    "تاريخ البداية", "تاريخ النهاية",
    "بداية خط الأساس", "نهاية خط الأساس",
    "البداية الفعلية", "النهاية الفعلية",
    "نسبة الإنجاز", "معلم",
]


def _norm_key(text):
    return normalize_light(text).lower()


class ScheduleImportWizard(models.TransientModel):
    _name = "rawasi.schedule.import.wizard"
    _description = "معالج استيراد الجدول الزمني"

    project_id = fields.Many2one("project.project", string="المشروع", required=True)
    file = fields.Binary(string="ملف الجدول الزمني (xlsx)", required=True)
    filename = fields.Char()
    replace_existing = fields.Boolean(string="استبدال الأنشطة الحالية", default=False)

    def _load_sheet(self):
        if openpyxl is None:
            raise UserError(_("مكتبة openpyxl غير متوفرة على الخادم."))
        try:
            data = base64.b64decode(self.file)
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise UserError(_("تعذّر قراءة الملف: %s") % exc)
        return wb.active

    def _detect_header(self, rows):
        alias_map = {_norm_key(k): v for k, v in HEADER_ALIASES.items()}
        best = None
        for r_idx, row in enumerate(rows[:HEADER_SCAN_ROWS]):
            mapping = {}
            for c_idx, cell in enumerate(row):
                field = alias_map.get(_norm_key(cell)) if cell else None
                if field and field not in mapping.values():
                    mapping[c_idx] = field
            logical = set(mapping.values())
            if "name" in logical and ({"date_start", "date_end"} & logical):
                score = len(logical)
                if best is None or score > best[0]:
                    best = (score, r_idx, mapping)
        if not best:
            raise UserError(
                _("الملف لا يطابق قالب جدول زمني: لم يُعثر على عمود «النشاط» وتواريخ.")
            )
        return best[1], best[2]

    @staticmethod
    def _to_date(value):
        if not value:
            return False
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(str(value).strip(), fmt).date()
            except ValueError:
                continue
        return False

    @staticmethod
    def _to_progress(value):
        if value in (None, ""):
            return 0.0
        try:
            return float(str(value).replace("%", "").strip())
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def _to_int(value, default):
        if value in (None, ""):
            return default
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return default

    # ── تحميل قالب الجدول الزمني (Excel جاهز للتعبئة) ────────────
    def _generate_schedule_template(self):
        if openpyxl is None:
            raise UserError(_("مكتبة openpyxl غير متوفرة على الخادم."))
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "الجدول الزمني"
        ws.append(TEMPLATE_COLUMNS)
        # صف مثال إرشادي (يمكن حذفه قبل الاستيراد)
        ws.append([
            "مثال: أعمال الحفر والردم", "الأعمال الإنشائية", 10,
            "2026-01-01", "2026-01-20",
            "2026-01-01", "2026-01-20",
            "", "", "0", "لا",
        ])
        for idx in range(1, len(TEMPLATE_COLUMNS) + 1):
            ws.column_dimensions[get_column_letter(idx)].width = 20
        # ورقة مرجعية بأسماء المراحل القياسية لنسخها في عمود «المرحلة»
        ws_ref = wb.create_sheet("المراحل المتاحة")
        ws_ref.append(["انسخ اسم المرحلة في عمود «المرحلة»"])
        for phase in self.env["rawasi.wbs.phase"].search([], order="sequence"):
            ws_ref.append([phase.name])
        ws_ref.column_dimensions["A"].width = 45
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def action_download_template(self):
        data = self._generate_schedule_template()
        attachment = self.env["ir.attachment"].create({
            "name": "قالب_الجدول_الزمني.xlsx",
            "type": "binary",
            "datas": base64.b64encode(data),
            "mimetype": (
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        })
        return {
            "type": "ir.actions.act_url",
            "url": "/web/content/%s?download=true" % attachment.id,
            "target": "self",
        }

    def action_import(self):
        self.ensure_one()
        sheet = self._load_sheet()
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise UserError(_("الملف فارغ."))
        header_idx, col_map = self._detect_header(rows)

        phases = self.env["rawasi.wbs.phase"].search([])
        phase_index = {_norm_key(p.name): p for p in phases}

        if self.replace_existing:
            self.project_id.wbs_activity_ids.unlink()

        Activity = self.env["rawasi.wbs.activity"]
        created = 0
        seq = 10
        for row in rows[header_idx + 1:]:
            data = {f: (row[c] if c < len(row) else None) for c, f in col_map.items()}
            name = data.get("name")
            if name in (None, ""):
                continue
            phase = phase_index.get(_norm_key(data.get("phase"))) if data.get("phase") else None
            is_milestone = bool(data.get("milestone")) and normalize_light(
                data.get("milestone")
            ).lower() in ("نعم", "yes", "true", "1", "معلم")
            Activity.create({
                "project_id": self.project_id.id,
                "phase_id": phase.id if phase else False,
                "name": str(name),
                "sequence": self._to_int(data.get("sequence"), seq),
                "date_start": self._to_date(data.get("date_start")),
                "date_end": self._to_date(data.get("date_end")),
                "baseline_start": self._to_date(data.get("baseline_start")),
                "baseline_end": self._to_date(data.get("baseline_end")),
                "actual_start": self._to_date(data.get("actual_start")),
                "actual_end": self._to_date(data.get("actual_end")),
                "progress": self._to_progress(data.get("progress")),
                "is_milestone": is_milestone,
            })
            created += 1
            seq += 10

        if not created:
            raise UserError(_("لم يتم العثور على أنشطة صالحة في الملف."))

        return {
            "type": "ir.actions.client",
            "tag": "display_notification",
            "params": {
                "title": _("اكتمل استيراد الجدول الزمني"),
                "message": _("تم استيراد %s نشاطاً.") % created,
                "type": "success",
                "sticky": False,
                "next": {"type": "ir.actions.act_window_close"},
            },
        }
