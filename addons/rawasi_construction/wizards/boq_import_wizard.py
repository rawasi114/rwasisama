# -*- coding: utf-8 -*-
"""الاستيراد الذكي لجدول الكميات من قوالب اعتماد (Header-Based, not Position-Based).

المنطق:
- يكتشف صف العناوين عبر مطابقة نصوص الأعمدة (بعد تطبيع عربي خفيف) بقاموس مرادفات.
- يرفض الملف إذا لم يجد عمودَي «الوصف» و«الكمية» (ليس جدول كميات).
- يطبّق Forward-Fill للفئة والمجموعة (الخلايا الفارغة ترث آخر قيمة).
- يطبّع رموز الوحدات ويطابقها بجدول الوحدات.
- يربط رمز SBC إن وُجد، ويسجّل تحذيرات (وحدة غير معروفة، كمية غير صالحة...).
"""
import base64
import io

from odoo import api, fields, models, _
from odoo.exceptions import UserError

from ..models.arabic_utils import normalize_light, normalize_unit

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    openpyxl = None

# مرادفات العناوين الخام -> الحقل المنطقي. تُطبَّع المفاتيح وقت التشغيل.
HEADER_ALIASES = {
    "الرقم التسلسلي": "serial", "الرقم": "serial", "مسلسل": "serial",
    "م": "serial", "ت": "serial",
    "الفئة": "category", "القسم": "category", "الباب": "category",
    "البند": "work_group", "المجموعة": "work_group", "البنود": "work_group",
    "وحدة القياس": "unit", "الوحدة": "unit", "وحدة": "unit",
    "الكمية": "quantity", "كمية": "quantity", "العدد": "quantity",
    "وصف البند": "description", "الوصف": "description", "وصف": "description",
    "البيان": "description", "بيان الأعمال": "description",
    "المواصفات": "specifications", "مواصفات": "specifications",
    "منتج من القائمة الإلزامية": "mandatory", "القائمة الإلزامية": "mandatory",
    "المحتوى المحلي": "mandatory", "إلزامي": "mandatory",
    "الرمز الإنشائي": "construction_code", "رمز sbc": "construction_code",
    "كود البناء": "construction_code", "الرمز": "construction_code",
    "مرفقات": "attachments", "المرفقات": "attachments",
    # الأسعار (لاستيراد جدول كميات مُسعّر مسبقاً)
    "سعر الوحدة": "unit_price", "السعر الافرادي": "unit_price",
    "السعر الإفرادي": "unit_price", "سعر البند": "unit_price",
    "السعر": "unit_price", "سعر": "unit_price",
    "إجمالي السعر": "total_price", "اجمالي السعر": "total_price",
    "الإجمالي": "total_price", "الاجمالي": "total_price", "القيمة": "total_price",
    "تكلفة الوحدة": "unit_cost", "سعر التكلفة": "unit_cost", "التكلفة": "unit_cost",
    "إجمالي التكلفة": "total_cost", "اجمالي التكلفة": "total_cost",
}

HEADER_SCAN_ROWS = 20  # عدد الصفوف الأولى التي نبحث فيها عن صف العناوين

# أعمدة قالب جدول الكميات (بالترتيب) — تطابق المرادفات أعلاه ليرفع الملف بسلاسة.
# عمود «تكلفة الوحدة» مُضاف لجمع التكاليف منذ مرحلة التسعير.
BOQ_TEMPLATE_COLUMNS = [
    "الرقم التسلسلي",
    "الفئة",
    "البند",
    "وصف البند",
    "المواصفات",
    "الوحدة",
    "الكمية",
    "تكلفة الوحدة",
    "سعر الوحدة",
    "إجمالي السعر",
    "منتج من القائمة الإلزامية",
    "الرمز الإنشائي",
]


def _norm_key(text):
    return normalize_light(text).lower()


class BoqImportWizard(models.TransientModel):
    _name = "rawasi.boq.import.wizard"
    _description = "معالج استيراد جدول الكميات"

    competition_id = fields.Many2one(
        "rawasi.competition", string="المنافسة", required=True
    )
    file = fields.Binary(string="ملف اعتماد (xlsx)", required=True)
    filename = fields.Char(string="اسم الملف")
    replace_existing = fields.Boolean(
        string="استبدال البنود الحالية",
        default=False,
        help="عند التفعيل تُحذف بنود المنافسة الحالية قبل الاستيراد.",
    )

    # ── قراءة الملف ──────────────────────────────────────────────
    def _load_sheet(self):
        if openpyxl is None:
            raise UserError(_("مكتبة openpyxl غير متوفرة على الخادم."))
        try:
            data = base64.b64decode(self.file)
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise UserError(_("تعذّر قراءة الملف كملف Excel صالح: %s") % exc)
        return wb.active

    @staticmethod
    def _normalized_alias_map():
        return {_norm_key(k): v for k, v in HEADER_ALIASES.items()}

    def _detect_header(self, rows):
        """يعيد (فهرس الصف، {رقم العمود: حقل}) لأفضل صف عناوين، أو يرفع خطأ."""
        alias_map = self._normalized_alias_map()
        best = None  # (score, row_index, mapping)
        for r_idx, row in enumerate(rows[:HEADER_SCAN_ROWS]):
            mapping = {}
            for c_idx, cell in enumerate(row):
                field = alias_map.get(_norm_key(cell)) if cell else None
                if field and field not in mapping.values():
                    mapping[c_idx] = field
            logical = set(mapping.values())
            # صالح فقط إذا توفّر الوصف والكمية معاً
            if {"description", "quantity"}.issubset(logical):
                score = len(logical)
                if best is None or score > best[0]:
                    best = (score, r_idx, mapping)
        if not best:
            raise UserError(
                _(
                    "الملف لا يطابق قالب جدول كميات اعتماد: لم يتم العثور على "
                    "عمودَي «وصف البند» و«الكمية». يرجى التأكد من الملف."
                )
            )
        return best[1], best[2]

    @staticmethod
    def _to_float(value):
        if value in (None, ""):
            return 0.0, False
        try:
            return float(value), True
        except (TypeError, ValueError):
            try:
                return float(str(value).replace(",", "").strip()), True
            except (TypeError, ValueError):
                return 0.0, False

    @staticmethod
    def _map_mandatory(value):
        token = normalize_light(value).lower() if value else ""
        if token in ("نعم", "yes"):
            return "yes"
        if token in ("لا", "no"):
            return "no"
        return False

    # ── تحميل قالب جدول الكميات (Excel جاهز للتعبئة) ─────────────
    def _generate_boq_template(self):
        if openpyxl is None:
            raise UserError(_("مكتبة openpyxl غير متوفرة على الخادم."))
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "جدول الكميات"
        ws.append(BOQ_TEMPLATE_COLUMNS)
        # صف مثال إرشادي (يمكن حذفه قبل الاستيراد)
        ws.append([
            1, "أعمال الخرسانة", "خرسانة مسلحة",
            "صبّ خرسانة C30 للأساسات",
            "خرسانة جاهزة مقاومتها 30 ميجاباسكال مع حديد تسليح حسب المخططات",
            "م3", 100, 320, 380, 38000, "نعم", "BC-100",
        ])
        for idx in range(1, len(BOQ_TEMPLATE_COLUMNS) + 1):
            ws.column_dimensions[get_column_letter(idx)].width = 22
        # ورقة مرجعية: الوحدات المتاحة في النظام (لنسخها في عمود «الوحدة»)
        ws_units = wb.create_sheet("الوحدات المتاحة")
        ws_units.append(["انسخ رمز الوحدة في عمود «الوحدة»"])
        for u in self.env["rawasi.unit"].search([("active", "=", True)], order="name"):
            ws_units.append([u.code or u.name, u.name or ""])
        ws_units.column_dimensions["A"].width = 14
        ws_units.column_dimensions["B"].width = 30
        # ورقة تعليمات قصيرة
        ws_help = wb.create_sheet("ملاحظات")
        ws_help.append(["الإلزامي: «وصف البند» و«الكمية» فقط."])
        ws_help.append(["باقي الأعمدة اختيارية — لكن «تكلفة الوحدة» و«سعر الوحدة» تجعل البند مُسعَّراً تلقائياً."])
        ws_help.append(["إن أدخلت «إجمالي السعر» دون «سعر الوحدة»، يُحتسب سعر الوحدة = الإجمالي ÷ الكمية."])
        ws_help.append(["الترتيب لا يهمّ — النظام يطابق أسماء الأعمدة."])
        ws_help.column_dimensions["A"].width = 90
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def action_download_template(self):
        data = self._generate_boq_template()
        attachment = self.env["ir.attachment"].create({
            "name": "قالب_جدول_الكميات.xlsx",
            "type": "binary",
            "datas": base64.b64encode(data),
            "mimetype": (
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        })
        return {
            "type": "ir.actions.act_url",
            "url": "/web/content/%s?download=true" % attachment.id,
            "target": "self",
        }

    # ── التنفيذ ──────────────────────────────────────────────────
    def action_import(self):
        self.ensure_one()
        comp = self.competition_id
        sheet = self._load_sheet()
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise UserError(_("الملف فارغ."))

        header_idx, col_map = self._detect_header(rows)
        unit_index = self.env["rawasi.unit"].build_alias_index()
        SbcCode = self.env["rawasi.sbc.code"]

        if self.replace_existing:
            comp.boq_item_ids.unlink()

        vals_list = []
        warnings = {"unit": 0, "qty": 0, "mandatory_blank": 0, "sbc": 0}
        priced_count = 0
        last_category = False
        last_group = False
        seq = 10

        for row in rows[header_idx + 1:]:
            data = {field: (row[c] if c < len(row) else None) for c, field in col_map.items()}
            desc = data.get("description")
            qty_raw = data.get("quantity")
            # تخطّي الصفوف الفارغة تماماً
            if (desc in (None, "")) and (qty_raw in (None, "")):
                continue
            if desc in (None, ""):
                # صف بلا وصف لكنه يحوي كمية -> تخطٍّ آمن
                continue

            # Forward-Fill للفئة والمجموعة
            category = data.get("category")
            category = category if category not in (None, "") else last_category
            last_category = category
            group = data.get("work_group")
            group = group if group not in (None, "") else last_group
            last_group = group

            qty, ok = self._to_float(qty_raw)
            if not ok:
                warnings["qty"] += 1

            unit_text = data.get("unit")
            unit = unit_index.get(normalize_unit(unit_text)) if unit_text else None
            if unit_text and not unit:
                warnings["unit"] += 1

            mandatory = self._map_mandatory(data.get("mandatory"))
            if data.get("mandatory") in (None, "") or (
                data.get("mandatory") and mandatory is False
            ):
                if data.get("mandatory") in (None, ""):
                    warnings["mandatory_blank"] += 1

            construction_code = data.get("construction_code")
            sbc = SbcCode.match_code(construction_code)
            if construction_code and str(construction_code).strip() not in ("", "غير محدد") and not sbc:
                warnings["sbc"] += 1

            # الأسعار: تُرفع كما لو سُعّر البند يدوياً في أودو. لو توفّر الإجمالي
            # فقط (دون سعر الوحدة) يُشتقّ سعر الوحدة = الإجمالي ÷ الكمية.
            unit_price, _up_ok = self._to_float(data.get("unit_price"))
            total_price, _tp_ok = self._to_float(data.get("total_price"))
            if not unit_price and total_price and qty:
                unit_price = total_price / qty
            unit_cost, _uc_ok = self._to_float(data.get("unit_cost"))
            total_cost, _tc_ok = self._to_float(data.get("total_cost"))
            if not unit_cost and total_cost and qty:
                unit_cost = total_cost / qty
            if unit_price:
                priced_count += 1

            vals_list.append({
                "competition_id": comp.id,
                "sequence": seq,
                "serial": str(data.get("serial")) if data.get("serial") not in (None, "") else False,
                "category": category or False,
                "work_group": group or False,
                "name": str(desc),
                "specifications": data.get("specifications") or False,
                "unit_text": str(unit_text) if unit_text else False,
                "unit_id": unit.id if unit else False,
                "quantity": qty,
                "mandatory_local": mandatory,
                "construction_code": str(construction_code) if construction_code not in (None, "") else False,
                "sbc_code_id": sbc.id if sbc else False,
                "unit_price": unit_price,
                "unit_cost": unit_cost,
            })
            seq += 10

        if not vals_list:
            raise UserError(_("لم يتم العثور على أي بنود صالحة في الملف."))

        self.env["rawasi.boq.item"].create(vals_list)
        if comp.state == "draft":
            comp.state = "pricing"

        summary = _(
            "تم استيراد %(count)s بنداً (منها %(priced)s مُسعّرة).\n"
            "تحذيرات: وحدات غير معروفة %(unit)s، كميات غير صالحة %(qty)s، "
            "محتوى محلي فارغ %(mand)s، رموز SBC غير مطابقة %(sbc)s."
        ) % {
            "count": len(vals_list),
            "priced": priced_count,
            "unit": warnings["unit"],
            "qty": warnings["qty"],
            "mand": warnings["mandatory_blank"],
            "sbc": warnings["sbc"],
        }
        comp.message_post(body=summary.replace("\n", "<br/>"))

        return {
            "type": "ir.actions.client",
            "tag": "display_notification",
            "params": {
                "title": _("اكتمل استيراد جدول الكميات"),
                "message": summary,
                "type": "success",
                "sticky": True,
                "next": {"type": "ir.actions.act_window_close"},
            },
        }
