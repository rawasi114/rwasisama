# -*- coding: utf-8 -*-
"""معالج استيراد جداول الكميات الذكي — Reference Items Import Wizard.

يستقبل ملف Excel + يربطه بمنافسة (linked_tender) إلزامياً + يحلّل صفوفه
بالمحلل الهرمي + يطبّق التوحيد التلقائي + يستخلص المواصفات + يشغّل محرك
المطابقة الذكية، ثم يفتح دفعة الاستيراد على شاشة المراجعة.
"""
import base64
import io

from odoo import api, fields, models, _
from odoo.exceptions import UserError

from odoo.addons.rawasi_construction.services.auto_normalization import (
    normalize_unit,
)
from odoo.addons.rawasi_construction.services.specifications_extractor import (
    extract_all as extract_specs,
)
from odoo.addons.rawasi_construction.services.matching_engine import find_match
from odoo.addons.rawasi_construction.services.hierarchy_classifier import classify


# ── الأعمدة المعتمدة في قالب اعتماد (وقالب رواسي الموسَّع) ────
_COL_ALIASES = {
    "serial":         ("الرقم التسلسلي", "الرقم", "م"),
    "main_category":  ("الفئة الرئيسية", "الفئة"),
    "item_group":     ("البند (المستوى الثاني)", "البند"),
    "simplified":     ("الوصف المبسط (المستوى الثالث)", "الوصف المبسط"),
    "full_desc":      ("الوصف الكامل كما بالجدول (المستوى الرابع)",
                       "وصف البند", "الوصف"),
    "unit":           ("وحدة القياس", "الوحدة"),
    "quantity":       ("الكمية",),
    "specifications": ("المواصفات",),
    "mandatory":      ("منتج من القائمة الإلزامية", "إلزامي"),
    "lcgpa":          ("الرمز الإنشائي", "LCGPA"),
    "unit_cost":      ("تكلفة الوحدة", "التكلفة", "Unit Cost", "Cost"),
    "unit_price":     ("سعر الوحدة", "السعر", "Unit Price", "Price"),
    "total_cost":     ("إجمالي التكلفة", "اجمالي التكلفة", "Total Cost"),
    "total_price":    ("إجمالي السعر", "اجمالي السعر", "Total Price"),
}


def _detect_columns(header_row):
    mapping = {}
    for idx, cell in enumerate(header_row):
        if not cell:
            continue
        clean = str(cell).strip()
        for key, aliases in _COL_ALIASES.items():
            if clean in aliases:
                mapping[key] = idx
                break
    return mapping


class ReferenceItemsImportWizard(models.TransientModel):
    _name = "rawasi.reference.items.import.wizard"
    _description = "معالج استيراد جدول كميات (سجل البنود المرجعي)"

    file_data = fields.Binary(string="ملف Excel", required=True)
    filename = fields.Char(string="اسم الملف", required=True)
    # ── الربط الإلزامي بالمنافسة ─────────────────────────────────
    linked_tender_id = fields.Many2one(
        "rawasi.competition", string="المنافسة المرتبطة",
        required=True,
        help="إلزامي. اختر منافسة قائمة أو أنشئ منافسة جديدة.",
    )
    source_entity_id = fields.Many2one(
        "res.partner", string="الجهة المصدر",
    )
    notes = fields.Text(string="ملاحظات")

    def action_analyze(self):
        """يقرأ الملف، يحلّله، ينشئ دفعة + سطورها، يفتح شاشة المراجعة."""
        self.ensure_one()
        if not self.linked_tender_id:
            raise UserError(_("يجب اختيار منافسة قبل تحليل الملف."))

        try:
            import openpyxl
        except ImportError:
            raise UserError(_("مكتبة openpyxl غير متوفّرة على الخادم."))

        try:
            wb = openpyxl.load_workbook(
                io.BytesIO(base64.b64decode(self.file_data)),
                data_only=True,
            )
        except Exception as e:
            raise UserError(_("تعذّر فتح الملف: %s") % e)

        # نقرأ كل الأوراق ونتعامل مع كل ورقة كقسم منفصل
        Batch = self.env["rawasi.import.batch"]
        BatchLine = self.env["rawasi.import.batch.line"]

        batch = Batch.create({
            "filename": self.filename,
            "file_data": self.file_data,
            "linked_tender_id": self.linked_tender_id.id,
            "source_entity_id": self.source_entity_id.id or False,
            "notes": self.notes or False,
            "state": "analyzing",
        })

        total_lines = 0
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            # نبحث عن صف العناوين
            header_idx = None
            for i, r in enumerate(rows):
                if r and any(
                    str(c).strip() in (_COL_ALIASES["full_desc"] + ("وصف البند",))
                    for c in r if c
                ):
                    header_idx = i
                    break
            if header_idx is None:
                continue

            col_map = _detect_columns(rows[header_idx])
            if "full_desc" not in col_map:
                continue

            last_mc, last_ig = None, None
            for r in rows[header_idx + 1:]:
                if not r or not any(r):
                    continue

                def _get(key):
                    idx = col_map.get(key)
                    if idx is None or idx >= len(r) or r[idx] is None:
                        return None
                    v = str(r[idx]).strip()
                    return v if v else None

                full_desc = _get("full_desc")
                if not full_desc:
                    continue
                total_lines += 1

                # ❶ المستويات الأربعة من قالب اعتماد + استدلال ذكي
                excel_mc = _get("main_category")
                excel_ig = _get("item_group")
                excel_sd = _get("simplified")
                # forward-fill
                if excel_mc:
                    last_mc = excel_mc
                if excel_ig:
                    last_ig = excel_ig
                # لو الـ Excel ما فيها هرمية، نستدل من النص الكامل
                inferred = classify(full_desc)
                mc = excel_mc or last_mc or inferred["main_category"]
                ig = excel_ig or last_ig or inferred["item_group"]
                sd = excel_sd or inferred["simplified_desc"]

                # ❷ التوحيد التلقائي للوحدات
                unit_raw = _get("unit") or ""
                unit_code, _label = normalize_unit(unit_raw)

                # ❸ استخلاص المواصفات بـ regex
                specs = extract_specs(full_desc)

                # ❹ كشف رمز LCGPA لو موجود في القاعدة
                lcgpa_raw = _get("lcgpa") or ""
                detected_lcgpa = False
                if lcgpa_raw and lcgpa_raw.strip().lower() not in ("غير محدد", "لا يوجد"):
                    lc = self.env["rawasi.lcgpa.code"].search(
                        [("code", "=", lcgpa_raw.strip())], limit=1,
                    )
                    if lc:
                        detected_lcgpa = lc.id

                # ❺ منظومة المطابقة الذكية
                ref, conf, source = find_match(
                    self.env, full_desc,
                    lcgpa_raw=lcgpa_raw, extracted_specs=specs,
                )

                # تحويل أعمدة التكلفة والسعر لأرقام (مع تجاهل الفواصل)
                def _to_float(key):
                    v = _get(key)
                    if not v:
                        return 0.0
                    try:
                        return float(v.replace(",", "").strip())
                    except (ValueError, AttributeError):
                        return 0.0

                BatchLine.create({
                    "import_batch_id": batch.id,
                    "line_number": total_lines,
                    "original_text": full_desc,
                    "original_unit": unit_raw or False,
                    "original_quantity": _get("quantity") or False,
                    "original_main_category": mc,
                    "original_item_group": ig,
                    "original_simplified_desc": sd,
                    "original_specifications": _get("specifications") or False,
                    "original_construction_code": lcgpa_raw or False,
                    "original_mandatory_local": _get("mandatory") or False,
                    "original_unit_cost": _to_float("unit_cost"),
                    "original_unit_price": _to_float("unit_price"),
                    "normalized_unit_code": unit_code or False,
                    "detected_lcgpa_id": detected_lcgpa,
                    "suggested_match_id": ref.id if ref else False,
                    "suggestion_confidence": conf,
                    "suggestion_source": source,
                    "extracted_dimensions": specs.get("dimensions") or False,
                    "extracted_thickness":  specs.get("thickness") or False,
                    "extracted_diameter":   specs.get("diameter") or False,
                    "extracted_strength":   specs.get("strength") or False,
                })

        if total_lines == 0:
            batch.unlink()
            raise UserError(_(
                "لم أتمكّن من قراءة أي سطر — تأكّد أن الملف يحتوي على "
                "عمود «الوصف الكامل كما بالجدول» أو «وصف البند»."
            ))

        batch.state = "reviewing"
        batch.message_post(body=_(
            "تم تحليل %d سطر — راجع/عدّل ثم اضغط «اعتماد الاستيراد»."
        ) % total_lines)

        return {
            "type": "ir.actions.act_window",
            "name": _("دفعة الاستيراد"),
            "res_model": "rawasi.import.batch",
            "res_id": batch.id,
            "view_mode": "form",
            "target": "current",
        }
